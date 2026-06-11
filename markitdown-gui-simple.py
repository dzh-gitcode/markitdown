import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

class DocumentConverter:
    def accepts(self, stream, stream_info):
        return True
    
    def convert(self, stream, stream_info):
        raise NotImplementedError

class PlainTextConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        content = stream.read().decode('utf-8', errors='replace')
        return type('Result', (), {'text_content': content})

class DocxConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import mammoth
            result = mammoth.convert_to_markdown(stream)
            return type('Result', (), {'text_content': result.value})
        except Exception as e:
            raise Exception(f"DOCX转换失败: {str(e)}")

class PdfConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import pdfplumber
            with pdfplumber.open(stream) as pdf:
                text = "\n\n".join([page.extract_text() for page in pdf.pages])
            return type('Result', (), {'text_content': text})
        except Exception as e:
            try:
                from pdfminer.high_level import extract_text
                import io
                text = extract_text(io.BytesIO(stream.read()))
                return type('Result', (), {'text_content': text})
            except Exception as e2:
                raise Exception(f"PDF转换失败: {str(e2)}")

class XlsxConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import pandas as pd
            df = pd.read_excel(stream)
            return type('Result', (), {'text_content': df.to_markdown()})
        except Exception as e:
            raise Exception(f"XLSX转换失败: {str(e)}")

class XlsConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import pandas as pd
            df = pd.read_excel(stream)
            return type('Result', (), {'text_content': df.to_markdown()})
        except Exception as e:
            raise Exception(f"XLS转换失败: {str(e)}")

class PptxConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            from pptx import Presentation
            prs = Presentation(stream)
            slides_content = []
            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        slide_text.append(shape.text)
                slides_content.append("\n".join(slide_text))
            return type('Result', (), {'text_content': "\n\n---\n\n".join(slides_content)})
        except Exception as e:
            raise Exception(f"PPTX转换失败: {str(e)}")

class HtmlConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            from markdownify import markdownify as md
            content = stream.read().decode('utf-8', errors='replace')
            return type('Result', (), {'text_content': md(content)})
        except Exception as e:
            raise Exception(f"HTML转换失败: {str(e)}")

class CsvConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import pandas as pd
            df = pd.read_csv(stream)
            return type('Result', (), {'text_content': df.to_markdown()})
        except Exception as e:
            content = stream.read().decode('utf-8', errors='replace')
            return type('Result', (), {'text_content': content})

class IpynbConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import json
            nb = json.load(stream)
            content = []
            for cell in nb.get('cells', []):
                if cell['cell_type'] == 'markdown':
                    content.append(''.join(cell['source']))
                elif cell['cell_type'] == 'code':
                    content.append(f"```python\n{''.join(cell['source'])}\n```")
            return type('Result', (), {'text_content': '\n\n'.join(content)})
        except Exception as e:
            raise Exception(f"IPYNB转换失败: {str(e)}")

class EpubConverter(DocumentConverter):
    def convert(self, stream, stream_info):
        try:
            import ebooklib
            from ebooklib import epub
            book = epub.read_epub(stream)
            content = []
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    content.append(item.get_content().decode('utf-8', errors='replace'))
            return type('Result', (), {'text_content': '\n\n'.join(content)})
        except Exception as e:
            raise Exception(f"EPUB转换失败: {str(e)}")

class MarkItDownConverter:
    def __init__(self):
        self.converters = {
            '.txt': PlainTextConverter(),
            '.md': PlainTextConverter(),
            '.docx': DocxConverter(),
            '.doc': DocxConverter(),
            '.pdf': PdfConverter(),
            '.xlsx': XlsxConverter(),
            '.xls': XlsConverter(),
            '.pptx': PptxConverter(),
            '.ppt': PptxConverter(),
            '.html': HtmlConverter(),
            '.htm': HtmlConverter(),
            '.csv': CsvConverter(),
            '.ipynb': IpynbConverter(),
            '.epub': EpubConverter(),
        }
    
    def convert(self, file_path):
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext in self.converters:
            converter = self.converters[file_ext]
            with open(file_path, 'rb') as fh:
                result = converter.convert(fh, None)
                return result
        else:
            raise Exception(f"不支持的文件格式: {file_ext}")

class MarkItDownGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("MarkItDown 文件转换器")
        self.root.geometry("700x500")
        self.root.resizable(True, True)
        
        # 初始化转换器
        self.converter = MarkItDownConverter()
        
        # 选中的文件列表
        self.selected_files = []
        
        # 创建主布局
        self.create_widgets()
        
    def create_widgets(self):
        # 顶部标题
        title_label = ttk.Label(self.root, text="MarkItDown 文件转换器", font=('Helvetica', 14, 'bold'))
        title_label.pack(pady=10)
        
        # 文件列表框
        self.file_list_frame = ttk.Frame(self.root)
        self.file_list_frame.pack(padx=10, pady=5, fill=tk.BOTH, expand=True)
        
        self.file_listbox = tk.Listbox(self.file_list_frame, selectmode=tk.MULTIPLE, width=80, height=15, font=('Helvetica', 10))
        self.file_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.scrollbar = ttk.Scrollbar(self.file_list_frame, orient=tk.VERTICAL, command=self.file_listbox.yview)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.file_listbox.config(yscrollcommand=self.scrollbar.set)
        
        # 按钮区域
        self.button_frame = ttk.Frame(self.root)
        self.button_frame.pack(pady=10)
        
        self.add_btn = ttk.Button(self.button_frame, text="添加文件", command=self.add_files)
        self.add_btn.grid(row=0, column=0, padx=5)
        
        self.remove_btn = ttk.Button(self.button_frame, text="移除选中", command=self.remove_selected)
        self.remove_btn.grid(row=0, column=1, padx=5)
        
        self.clear_btn = ttk.Button(self.button_frame, text="清空列表", command=self.clear_list)
        self.clear_btn.grid(row=0, column=2, padx=5)
        
        # 输出目录选择
        self.output_frame = ttk.Frame(self.root)
        self.output_frame.pack(padx=10, pady=5, fill=tk.X)
        
        ttk.Label(self.output_frame, text="输出目录:").pack(side=tk.LEFT)
        self.output_path_var = tk.StringVar(value=os.path.expanduser("~\\Desktop"))
        self.output_entry = ttk.Entry(self.output_frame, textvariable=self.output_path_var, width=60)
        self.output_entry.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        self.browse_output_btn = ttk.Button(self.output_frame, text="浏览", command=self.browse_output)
        self.browse_output_btn.pack(side=tk.RIGHT, padx=5)
        
        # 转换按钮和进度
        self.convert_btn = ttk.Button(self.root, text="开始转换", command=self.convert_files, style='Accent.TButton')
        self.convert_btn.pack(pady=10)
        
        # 进度条
        self.progress = ttk.Progressbar(self.root, orient=tk.HORIZONTAL, length=600, mode='determinate')
        self.progress.pack(pady=5)
        
        # 状态标签
        self.status_var = tk.StringVar(value="就绪")
        self.status_label = ttk.Label(self.root, textvariable=self.status_var, foreground="blue")
        self.status_label.pack(pady=5)
        
    def add_files(self):
        files = filedialog.askopenfilenames(
            title="选择文件",
            filetypes=[
                ("所有支持的文件", "*.pdf;*.docx;*.doc;*.xlsx;*.xls;*.pptx;*.ppt;*.txt;*.html;*.md;*.csv;*.ipynb;*.epub"),
                ("PDF文件", "*.pdf"),
                ("Word文档", "*.docx;*.doc"),
                ("Excel表格", "*.xlsx;*.xls"),
                ("PowerPoint演示", "*.pptx;*.ppt"),
                ("文本文件", "*.txt"),
                ("HTML文件", "*.html"),
                ("Jupyter笔记本", "*.ipynb"),
                ("EPUB电子书", "*.epub"),
                ("所有文件", "*.*")
            ]
        )
        
        for file in files:
            if file not in self.selected_files:
                self.selected_files.append(file)
                self.file_listbox.insert(tk.END, os.path.basename(file))
        
        self.update_status(f"已添加 {len(files)} 个文件")
    
    def remove_selected(self):
        selected_indices = self.file_listbox.curselection()
        for index in reversed(selected_indices):
            del self.selected_files[index]
            self.file_listbox.delete(index)
        
        self.update_status(f"已移除 {len(selected_indices)} 个文件")
    
    def clear_list(self):
        self.selected_files.clear()
        self.file_listbox.delete(0, tk.END)
        self.update_status("列表已清空")
    
    def browse_output(self):
        path = filedialog.askdirectory(title="选择输出目录")
        if path:
            self.output_path_var.set(path)
    
    def update_status(self, message):
        self.status_var.set(message)
        self.root.update_idletasks()
    
    def convert_files(self):
        if not self.selected_files:
            messagebox.showwarning("警告", "请先添加要转换的文件")
            return
        
        output_dir = self.output_path_var.get()
        if not os.path.isdir(output_dir):
            messagebox.showwarning("警告", "输出目录无效")
            return
        
        total_files = len(self.selected_files)
        success_count = 0
        fail_count = 0
        failed_files = []
        
        self.progress['maximum'] = total_files
        self.progress['value'] = 0
        
        for i, file_path in enumerate(self.selected_files):
            try:
                self.update_status(f"正在转换: {os.path.basename(file_path)}")
                
                # 调用转换器
                result = self.converter.convert(file_path)
                
                # 生成输出文件名
                base_name = os.path.splitext(os.path.basename(file_path))[0]
                output_path = os.path.join(output_dir, f"{base_name}.md")
                
                # 写入MD文件
                with open(output_path, 'w', encoding='utf-8') as f:
                    if hasattr(result, 'text_content'):
                        f.write(result.text_content)
                    elif isinstance(result, str):
                        f.write(result)
                    else:
                        f.write(str(result))
                
                success_count += 1
                
            except Exception as e:
                fail_count += 1
                failed_files.append(f"{os.path.basename(file_path)}: {str(e)}")
            
            self.progress['value'] = i + 1
            self.root.update_idletasks()
        
        self.update_status(f"转换完成: 成功 {success_count} 个, 失败 {fail_count} 个")
        
        if failed_files:
            error_message = "\n".join(failed_files)
            messagebox.showerror("转换失败", f"以下文件转换失败:\n{error_message}")
        else:
            messagebox.showinfo("转换完成", f"所有 {success_count} 个文件转换成功！")

def main():
    root = tk.Tk()
    app = MarkItDownGUI(root)
    root.mainloop()

if __name__ == "__main__":
    main()